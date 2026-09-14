import base64
import hashlib
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from odoo import _, fields, models
from odoo.exceptions import UserError, ValidationError


REQUIRED_SHEETS = {
    "CONFIG",
    "SERVICIOS",
    "TARIFAS",
    "RECARGOS",
    "LIMITES",
    "ZONAS",
    "PROVINCIAS_LIMITROFES",
    "CP_PROVINCIAS",
}


def _clean(value):
    return str(value or "").strip()


def _bool(value):
    if isinstance(value, bool):
        return value
    return _clean(value).lower() in {"1", "true", "yes", "si", "sí", "x"}


def _date(value):
    if not value:
        return False
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValidationError(_("Fecha no válida en plantilla GLS: %s") % text)


def _float(value, default=0.0):
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean(value).replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    return float(text)


def _rows(sheet):
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_clean(value) for value in values[0]]
    result = []
    for row in values[1:]:
        if not any(value not in (None, "") for value in row):
            continue
        result.append({headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))})
    return result


class OptimaGlsTariffImportWizard(models.TransientModel):
    _name = "optima.gls.tariff.import.wizard"
    _description = "Import Optima GLS Tariff Template"

    file_data = fields.Binary(string="GLS tariff Excel", required=True)
    filename = fields.Char(required=True)
    enable_matching_carriers = fields.Boolean(
        string="Activar motor de tarifas en transportistas GLS compatibles",
        default=True,
    )

    def action_import(self):
        self.ensure_one()
        try:
            payload = base64.b64decode(self.file_data or b"")
            workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)
        except Exception as exc:
            raise UserError(_("No se puede abrir la plantilla XLSX de GLS: %s") % exc) from exc

        missing = REQUIRED_SHEETS - set(workbook.sheetnames)
        if missing:
            raise ValidationError(_("Faltan hojas obligatorias en la plantilla GLS: %s") % ", ".join(sorted(missing)))

        config_rows = _rows(workbook["CONFIG"])
        # CONFIG is intentionally key/value, with its header on row 3.
        if not config_rows:
            values = list(workbook["CONFIG"].iter_rows(min_row=4, values_only=True))
            config = {_clean(row[0]): row[1] for row in values if row and _clean(row[0])}
        else:
            # The generic parser sees decorative rows, so use the stable cells.
            values = list(workbook["CONFIG"].iter_rows(min_row=4, values_only=True))
            config = {_clean(row[0]): row[1] for row in values if row and _clean(row[0])}

        book_code = _clean(config.get("BOOK_CODE"))
        book_name = _clean(config.get("BOOK_NAME")) or book_code
        template_version = _clean(config.get("TEMPLATE_VERSION"))
        currency_code = _clean(config.get("CURRENCY")) or "EUR"
        if not book_code:
            raise ValidationError(_("CONFIG/BOOK_CODE es obligatorio."))
        currency = self.env["res.currency"].search([("name", "=", currency_code)], limit=1)
        if not currency:
            raise ValidationError(_("No existe la moneda %s en Odoo.") % currency_code)

        book_model = self.env["optima.gls.tariff.book"].sudo()
        book = book_model.search([("code", "=", book_code)], limit=1)
        if book:
            book.service_ids.unlink()
            book.postal_ids.unlink()
            book.border_ids.unlink()
            book.write(
                {
                    "name": book_name,
                    "template_version": template_version,
                    "currency_id": currency.id,
                    "source_filename": self.filename,
                    "source_hash": hashlib.sha256(payload).hexdigest(),
                    "imported_at": fields.Datetime.now(),
                    "active": True,
                }
            )
        else:
            book = book_model.create(
                {
                    "code": book_code,
                    "name": book_name,
                    "template_version": template_version,
                    "currency_id": currency.id,
                    "source_filename": self.filename,
                    "source_hash": hashlib.sha256(payload).hexdigest(),
                    "imported_at": fields.Datetime.now(),
                    "active": True,
                }
            )

        postal_model = self.env["optima.gls.tariff.postal"].sudo()
        for row in _rows(workbook["CP_PROVINCIAS"]):
            prefix = _clean(row.get("postal_prefix"))
            province = _clean(row.get("province"))
            if prefix and province:
                postal_model.create(
                    {
                        "book_id": book.id,
                        "postal_prefix": prefix.zfill(2),
                        "province": province,
                        "source": _clean(row.get("source")),
                    }
                )

        border_model = self.env["optima.gls.tariff.border"].sudo()
        for row in _rows(workbook["PROVINCIAS_LIMITROFES"]):
            origin = _clean(row.get("origin_province"))
            destination = _clean(row.get("destination_province"))
            if origin and destination:
                border_model.create(
                    {
                        "book_id": book.id,
                        "origin_province": origin,
                        "destination_province": destination,
                        "source": _clean(row.get("source")),
                        "confirmed_by_gls": _bool(row.get("confirmed_by_gls")),
                        "notes": _clean(row.get("notes")),
                    }
                )

        service_model = self.env["optima.gls.tariff.service"].sudo()
        services = {}
        for row in _rows(workbook["SERVICIOS"]):
            code = _clean(row.get("service_code"))
            if not code:
                continue
            valid_from = _date(row.get("valid_from"))
            valid_to = _date(row.get("valid_to"))
            if not valid_from or not valid_to:
                raise ValidationError(_("El servicio %s necesita valid_from y valid_to.") % code)
            service = service_model.create(
                {
                    "book_id": book.id,
                    "code": code,
                    "name": _clean(row.get("name")) or code,
                    "valid_from": valid_from,
                    "valid_to": valid_to,
                    "gls_service_hint": _clean(row.get("gls_service_hint")),
                    "gls_shiptime_hint": _clean(row.get("gls_shiptime_hint")),
                    "volumetric_factor_kg_m3": _float(row.get("volumetric_factor_kg_m3"), 167.0),
                    "mono_parcel": _bool(row.get("mono_parcel")),
                    "max_packages": int(_float(row.get("max_packages"), 1.0)),
                    "zone_mode": _clean(row.get("zone_mode")) or "EXPLICIT_RULES",
                    "origin_postal_code": _clean(row.get("origin_postal_code")),
                    "origin_province": _clean(row.get("origin_province")),
                    "additional_kg_rounding": _clean(row.get("additional_kg_rounding")) or "CEIL",
                    "surcharge_mode": _clean(row.get("surcharge_mode")) or "ADD_ON_BASE",
                    "active": _bool(row.get("active")),
                    "notes": _clean(row.get("notes")),
                }
            )
            services[code] = service

        rate_model = self.env["optima.gls.tariff.rate"].sudo()
        for row in _rows(workbook["TARIFAS"]):
            service = services.get(_clean(row.get("service_code")))
            if not service:
                continue
            zone = _clean(row.get("zone_code"))
            if not zone:
                continue
            rate_model.create(
                {
                    "service_id": service.id,
                    "zone_code": zone,
                    "weight_to_kg": _float(row.get("weight_to_kg")),
                    "price": _float(row.get("price_eur")),
                    "additional_kg_price": _float(row.get("additional_kg_price_eur")),
                    "notes": _clean(row.get("notes")),
                }
            )

        surcharge_model = self.env["optima.gls.tariff.surcharge"].sudo()
        for row in _rows(workbook["RECARGOS"]):
            service = services.get(_clean(row.get("service_code")))
            if not service:
                continue
            code = _clean(row.get("surcharge_code"))
            if not code:
                continue
            surcharge_model.create(
                {
                    "service_id": service.id,
                    "code": code,
                    "kind": _clean(row.get("kind")) or "PERCENT",
                    "value": _float(row.get("value")),
                    "apply_mode": _clean(row.get("apply_mode")) or "ADD_ON_BASE",
                    "valid_from": _date(row.get("valid_from")),
                    "valid_to": _date(row.get("valid_to")),
                    "active": _bool(row.get("active")),
                    "notes": _clean(row.get("notes")),
                }
            )

        limit_model = self.env["optima.gls.tariff.limit"].sudo()
        for row in _rows(workbook["LIMITES"]):
            service = services.get(_clean(row.get("service_code")))
            if not service:
                continue
            package_type = _clean(row.get("package_type"))
            if not package_type:
                continue
            limit_model.create(
                {
                    "service_id": service.id,
                    "package_type": package_type,
                    "max_weight_kg": _float(row.get("max_weight_kg")),
                    "max_sum_sides_cm": _float(row.get("max_sum_sides_cm")),
                    "max_length_girth_cm": _float(row.get("max_length_girth_cm")),
                    "max_side_cm": _float(row.get("max_side_cm")),
                    "enforced_by_default": _bool(row.get("enforced_by_default")),
                    "notes": _clean(row.get("notes")),
                }
            )

        zone_model = self.env["optima.gls.tariff.zone"].sudo()
        for row in _rows(workbook["ZONAS"]):
            service = services.get(_clean(row.get("service_code")))
            if not service:
                continue
            zone_code = _clean(row.get("zone_code"))
            rule_type = _clean(row.get("rule_type"))
            if not zone_code or not rule_type:
                continue
            zone_model.create(
                {
                    "service_id": service.id,
                    "sequence": int(_float(row.get("sequence"), 100.0)),
                    "rule_type": rule_type,
                    "country_code": _clean(row.get("country_code")).upper(),
                    "postal_prefix": _clean(row.get("postal_prefix")),
                    "origin_province": _clean(row.get("origin_province")),
                    "destination_province": _clean(row.get("destination_province")),
                    "zone_code": zone_code,
                    "notes": _clean(row.get("notes")),
                }
            )

        if self.enable_matching_carriers:
            carriers = self.env["delivery.carrier"].sudo().search([("delivery_type", "=", "gls_asm")])
            for carrier in carriers:
                inferred = carrier._optima_gls_tariff_service_code()
                if inferred and inferred in services:
                    carrier.write({"optima_gls_tariff_enabled": True})

        return {
            "type": "ir.actions.act_window",
            "res_model": "optima.gls.tariff.book",
            "res_id": book.id,
            "view_mode": "form",
            "target": "current",
        }
